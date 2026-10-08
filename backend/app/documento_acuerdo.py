"""
Documento Word del acuerdo de pago / avenimiento.

Replica el formato de los acuerdos que se envían al mandante para su
aprobación (calculadora `cobra369`), pero con los datos de cada
organización: membrete, razón social, datos de transferencia (los del
mandante si tiene propios, si no los del estudio) y texto de cierre.

  - Encabezado de la primera hoja: membrete a la izquierda; N° de cobranza,
    fecha y filial a la derecha. Hojas siguientes: solo el membrete.
  - Párrafos de derivación (con o sin abonos previos al acuerdo).
  - Total del pagaré y texto de las cuotas.
  - Datos de transferencia en ficha.
  - Tabla de cuotas: con desglose (capital, interés, honorarios, gastos,
    comisión) si el acuerdo se armó con la calculadora; si no, solo monto.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import _Cell

from app.calculos import MESES, clp
from app.models.acuerdo import AcuerdoPago, Cuota
from app.models.cobranza import Cobranza
from app.models.empresa import Empresa
from app.rut import rut_con_puntos

FUENTE = "Times New Roman"
CIERRE_DEFECTO = (
    "SOLICITAMOS se sirva tener por presentado el presente acuerdo, prestarle su "
    "aprobación y redactar el documento que debe firmar el deudor."
)

CERO = Decimal(0)


def _miles(v) -> str:
    return f"{int(round(Decimal(v or 0))):,}".replace(",", ".")


def _fecha_larga(f: date) -> str:
    return f"{f.day:02d} {MESES[f.month - 1]} {f.year}"


def _fecha_texto(f: date) -> str:
    return f"{f.day} de {MESES[f.month - 1]} de {f.year}"


def _fecha_tabla(f: Optional[date]) -> str:
    return f.strftime("%d-%m-%y") if f else ""


def _parrafo(doc, texto: str = "", *, negrita=False, centrado=False, tam=12,
             subrayado=False, justificado=True, antes=0, despues=14):
    p = doc.add_paragraph()
    p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if centrado
                   else WD_ALIGN_PARAGRAPH.JUSTIFY if justificado else WD_ALIGN_PARAGRAPH.LEFT)
    f = p.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(antes), Pt(despues), 1.5
    if texto:
        r = p.add_run(texto)
        r.bold, r.underline = negrita, subrayado
        r.font.name, r.font.size = FUENTE, Pt(tam)
    return p


def _sombrear(celda, color_hex: str) -> None:
    tc_pr = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color_hex)
    tc_pr.append(sombra)


def _celda(celda, texto: str, *, negrita=False, tam=11) -> None:
    celda.text = ""
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(texto)
    r.bold = negrita
    r.font.name, r.font.size = FUENTE, Pt(tam)


def _encabezados(doc, emp: Empresa, cob: Cobranza, hoy: date) -> None:
    seccion = doc.sections[0]
    seccion.different_first_page_header_footer = True

    def membrete(contenedor):
        insertar_logo(contenedor, emp)

    # Primera hoja: membrete + datos del caso a la derecha.
    tabla = seccion.first_page_header.add_table(rows=1, cols=2, width=Cm(16.6))
    izq, der = tabla.rows[0].cells
    membrete(izq)
    der.text = ""
    filial = (cob.filial.nombre if cob.filial else
              (cob.cliente.nombre_fantasia or cob.cliente.razon_social) if cob.cliente else "")
    for i, linea in enumerate((f"COB. {cob.numero}", _fecha_larga(hoy),
                               f"FILIAL {filial.upper()}" if filial else "")):
        p = der.paragraphs[0] if i == 0 else der.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(linea)
        r.font.name, r.font.size = FUENTE, Pt(12)

    # Hojas siguientes: solo el membrete.
    membrete(seccion.header)


def insertar_logo(contenedor, emp: Empresa, alto_cm: float = 2.2) -> None:
    """
    Logo de la empresa en el encabezado. Si todavía no lo cargó, queda un
    recuadro que marca dónde va (sin ningún nombre: cada empresa pone el suyo).
    """
    p = contenedor.paragraphs[0]
    if emp.logo:
        p.add_run().add_picture(BytesIO(emp.logo), height=Cm(alto_cm))
        return
    if isinstance(contenedor, _Cell):
        tabla = contenedor.add_table(rows=1, cols=1)
    else:
        tabla = contenedor.add_table(rows=1, cols=1, width=Cm(4.5))
    celda = tabla.rows[0].cells[0]
    celda.width = Cm(4.5)
    _sombrear(celda, "EEEEEE")
    _celda(celda, "LOGO DE LA EMPRESA", tam=8)
    celda.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    tabla.rows[0].height = Cm(alto_cm)


def _texto_cuotas(acuerdo: AcuerdoPago, cuotas: List[Cuota]) -> str:
    partes = []
    pie = Decimal(acuerdo.pie or 0)
    if pie > 0:
        partes.append(f"Se efectúa un abono inicial (PIE) de {clp(pie)}.")
    sujeto = "El saldo restante" if pie > 0 else "La deuda"
    n = len(cuotas)
    valor = cuotas[0].monto
    iguales = all(c.monto == valor for c in cuotas)
    primera, ultima = cuotas[0].fecha_vencimiento, cuotas[-1].fecha_vencimiento
    if n == 1:
        partes.append(f"{sujeto} se pagará en 1 cuota de {clp(valor)}, con vencimiento el "
                      f"{_fecha_texto(primera)}.")
    elif iguales:
        dia = acuerdo.dia_pago or primera.day
        partes.append(
            f"{sujeto} se pagará en {n} cuotas iguales, mensuales y sucesivas de {clp(valor)} "
            f"cada una. La primera cuota vence el {_fecha_texto(primera)} y las restantes los "
            f"días {dia} de cada mes, finalizando el {_fecha_texto(ultima)}, ambas fechas inclusive."
        )
    else:
        partes.append(f"{sujeto} se pagará en {n} cuotas mensuales según el detalle siguiente, "
                      f"desde el {_fecha_texto(primera)} hasta el {_fecha_texto(ultima)}.")
    return " ".join(partes)


def _bloque_pago(doc, instrucciones: str) -> None:
    lineas = [l.strip() for l in (instrucciones or "").splitlines() if l.strip()]
    if not lineas:
        return
    _parrafo(doc, "El pago deberá realizarse de la siguiente forma:", despues=6)
    tabla = doc.add_table(rows=1, cols=1)
    tabla.alignment = WD_TABLE_ALIGNMENT.LEFT
    celda = tabla.rows[0].cells[0]
    _sombrear(celda, "F2F2F2")
    celda.text = ""
    for i, linea in enumerate(lineas):
        p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        etiqueta, sep, valor = linea.partition(":")
        if sep and len(etiqueta) <= 30:
            r = p.add_run(etiqueta + ":")
            r.bold = True
            r.font.name, r.font.size = FUENTE, Pt(11)
            r = p.add_run(" " + valor.strip())
        else:
            r = p.add_run(linea)
        r.font.name, r.font.size = FUENTE, Pt(11)
    _parrafo(doc, despues=4)


def _tabla_cuotas(doc, acuerdo: AcuerdoPago, cuotas: List[Cuota], judicial: bool,
                  nombre_mandante: str) -> None:
    con_desglose = all(c.capital is not None for c in cuotas)
    pie = Decimal(acuerdo.pie or 0)
    if con_desglose:
        con_gastos = any(Decimal(c.gastos_judiciales or 0) > 0 for c in cuotas)
        con_comision = any(Decimal(c.comision or 0) > 0 for c in cuotas)
        cab = ["N° CUOTA", "FECHA", f"ABONO {nombre_mandante.upper()}"[:28], "INTERÉS",
               "HONORARIOS" if judicial else "GASTO COBRANZA"]
        cab += (["GASTOS JUDICIALES"] if con_gastos else []) + (["COMISIÓN"] if con_comision else [])
        cab += ["TOTAL CUOTA"]
    else:
        con_gastos = con_comision = False
        cab = ["N° CUOTA", "FECHA", "TOTAL CUOTA"]

    tabla = doc.add_table(rows=1, cols=len(cab))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, t in enumerate(cab):
        _celda(tabla.rows[0].cells[i], t, negrita=True, tam=8.5)
        _sombrear(tabla.rows[0].cells[i], "E9E9E9")

    def fila(valores, negrita_total=True):
        celdas = tabla.add_row().cells
        for i, v in enumerate(valores):
            _celda(celdas[i], v, negrita=(negrita_total and i == len(valores) - 1) or i == 0 and v == "PIE")

    if pie > 0:
        if con_desglose:
            hon_pie = Decimal(acuerdo.honorarios or 0) - sum((Decimal(c.honorarios or 0) for c in cuotas), CERO)
            cap_pie = pie - hon_pie
            fila(["PIE", "", _miles(cap_pie), "0", _miles(hon_pie)]
                 + (["0"] if con_gastos else []) + (["0"] if con_comision else []) + [_miles(pie)])
        else:
            fila(["PIE", "", _miles(pie)])
    for c in cuotas:
        if con_desglose:
            fila([str(c.numero_cuota), _fecha_tabla(c.fecha_vencimiento), _miles(c.capital),
                  _miles(c.intereses), _miles(c.honorarios)]
                 + ([_miles(c.gastos_judiciales)] if con_gastos else [])
                 + ([_miles(c.comision)] if con_comision else []) + [_miles(c.monto)])
        else:
            fila([str(c.numero_cuota), _fecha_tabla(c.fecha_vencimiento), _miles(c.monto)])

    total = tabla.add_row().cells
    unidas = total[0].merge(total[len(cab) - 2])
    _celda(unidas, "TOTAL", negrita=True)
    _celda(total[len(cab) - 1], _miles(acuerdo.monto_total_acordado), negrita=True)


def generar(acuerdo: AcuerdoPago, cob: Cobranza, emp: Empresa, pagado_antes: Decimal,
            texto_cierre: Optional[str], hoy: date) -> BytesIO:
    judicial = cob.tipo == "judicial"
    deudor = cob.deudor
    cliente = cob.cliente
    nombre = (deudor.nombre if deudor else "_______________").upper()
    ident = " ".join(x for x in (
        f"ID {cob.id_externo}" if cob.id_externo else "",
        f"RUT {rut_con_puntos(deudor.rut)}" if deudor and deudor.rut else "",
    ) if x)
    empresa = emp.razon_social
    cuotas = list(acuerdo.cuotas)

    doc = Document()
    seccion = doc.sections[0]
    seccion.page_height, seccion.page_width = Cm(29.7), Cm(21.0)
    seccion.top_margin, seccion.bottom_margin = Cm(4.0), Cm(2.0)
    seccion.left_margin = seccion.right_margin = Cm(2.2)
    estilo = doc.styles["Normal"]
    estilo.font.name, estilo.font.size = FUENTE, Pt(12)

    _encabezados(doc, emp, cob, hoy)

    _parrafo(doc, "AVENIMIENTO" if judicial else "ACUERDO DE PAGO", negrita=True,
             subrayado=True, centrado=True, tam=14, despues=24)
    derivado = Decimal(cob.monto_original)
    if pagado_antes > 0:
        _parrafo(doc, f"{nombre} {ident} pasa a cobranza de {empresa} por un monto de "
                      f"${_miles(derivado)} y, luego de abonos por ${_miles(pagado_antes)}, queda un "
                      f"saldo remanente de ${_miles(derivado - pagado_antes)}.")
        _parrafo(doc, f"De acuerdo a este saldo pendiente las partes acuerdan celebrar el siguiente "
                      f"acuerdo, por el cual {nombre} se obliga a pagar la deuda antes descrita de la "
                      f"siguiente forma:")
    else:
        _parrafo(doc, f"Se informa que la cuenta de {nombre} {ident} ha sido derivada para su "
                      f"regularización a {empresa} por el monto de ${_miles(derivado)}.")
        _parrafo(doc, f"Sobre este monto pendiente las partes acuerdan celebrar el siguiente "
                      f"acuerdo, por el cual {nombre} se obliga a pagar la deuda antes descrita de la "
                      f"siguiente forma:")

    _parrafo(doc, "PAGARÉ EN CUOTAS A REALIZAR", negrita=True, centrado=True, antes=10, despues=0)
    _parrafo(doc, f"TOTAL PAGARÉ: ${_miles(acuerdo.monto_total_acordado)}", negrita=True,
             centrado=True, despues=18)
    _parrafo(doc, _texto_cuotas(acuerdo, cuotas))

    instrucciones = ((cliente.instrucciones_pago if cliente else None)
                     or emp.instrucciones_pago or "")
    _bloque_pago(doc, instrucciones)

    nombre_mandante = (cliente.nombre_fantasia or cliente.razon_social) if cliente else "mandante"
    _tabla_cuotas(doc, acuerdo, cuotas, judicial, nombre_mandante)

    _parrafo(doc, texto_cierre or CIERRE_DEFECTO, antes=20)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
