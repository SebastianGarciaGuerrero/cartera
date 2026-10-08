"""
Fase 2: calculadora 3-6-9 y acuerdos asistidos, UF del día, agenda,
recordatorios y mensaje de pago.
"""

from datetime import date, timedelta
from decimal import Decimal

from app import calculos
from app.indicadores import hoy_chile
from tests.conftest import crear_cartera, crear_organizacion, agregar_usuario, login

UF = Decimal("39000")


# ------------------------------------------------------------ cálculos puros

def test_honorarios_369_por_tramos():
    # 1.000.000 con UF 39.000: 10 UF = 390.000 al 9 % + 610.000 al 6 %.
    h = calculos.honorarios(Decimal(1_000_000), UF)
    assert calculos.redondear(h.total_honorarios) == 35_100 + 36_600
    # Sobre 50 UF entra el tramo del 3 %.
    h = calculos.honorarios(Decimal(3_000_000), UF)
    esperado = 390_000 * 0.09 + 1_560_000 * 0.06 + (3_000_000 - 1_950_000) * 0.03
    assert calculos.redondear(h.total_honorarios) == round(esperado)


def test_abono_se_separa_exacto():
    h = calculos.capital_desde_abono(Decimal(1_071_700), UF)
    assert calculos.redondear(h.capital) == 1_000_000
    j = calculos.capital_desde_abono(Decimal(110_000), None, "judicial")
    assert calculos.redondear(j.capital) == 100_000


def test_plan_cuadra_el_capital_y_redondea():
    plan = calculos.calcular_acuerdo(
        Decimal(1_234_567), 7, Decimal("1.5"), UF, abono_inicial=Decimal(200_000),
        redondeo="arriba", fecha_primera=date(2026, 1, 31),
    )
    assert plan.capital_pie + sum(f.capital for f in plan.cuotas) == 1_234_567
    assert plan.valor_cuota % 1000 == 0
    assert all(f.total == plan.valor_cuota for f in plan.cuotas)
    assert [f.fecha for f in plan.cuotas][:3] == [date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 31)]
    assert plan.gran_total == plan.valor_cuota * 7 + 200_000
    assert "7 cuotas iguales" in calculos.texto_acuerdo(plan)


# ------------------------------------------------------------ API

def test_calculadora_es_premium(cliente_http):
    base = crear_organizacion(plan="base")
    h = login(cliente_http, base["email"])
    r = cliente_http.post("/api/calculadora/honorarios", headers=h,
                          json={"capital": 1000000, "uf": 39000})
    assert r.status_code == 402


def test_acuerdo_asistido_crea_cuotas_con_desglose(cliente_http, org_a):
    h = login(cliente_http, org_a["email"])
    cartera = crear_cartera(cliente_http, h, monto=1_000_000, n=40)
    cuerpo = {
        "cobranza_id": cartera["cobranza_id"], "capital": 1000000, "numero_cuotas": 4,
        "tasa_mensual": 1, "uf": 39000, "abono_inicial": 100000,
        "fecha_primera_cuota": (date.today() + timedelta(days=10)).isoformat(),
        "redondeo": "arriba",
    }
    simulado = cliente_http.post("/api/calculadora/acuerdo", headers=h, json=cuerpo).json()
    r = cliente_http.post("/api/calculadora/acuerdo/crear", headers=h, json=cuerpo)
    assert r.status_code == 201, r.text
    acuerdo = r.json()
    assert len(acuerdo["cuotas"]) == 4
    assert float(acuerdo["monto_total_acordado"]) == float(simulado["gran_total"])
    cuota = acuerdo["cuotas"][0]
    assert float(cuota["capital"]) + float(cuota["intereses"]) + float(cuota["honorarios"]) == float(cuota["monto"])

    cob = cliente_http.get(f"/api/cobranzas/{cartera['cobranza_id']}", headers=h).json()
    assert cob["estado"] == "acuerdo_pago"
    historial = cliente_http.get("/api/gestiones/", params={"cobranza_id": cartera["cobranza_id"]}, headers=h).json()
    assert any("ACUERDO DE PAGO" in g["descripcion"] for g in historial)
    # Un solo acuerdo vigente.
    assert cliente_http.post("/api/calculadora/acuerdo/crear", headers=h, json=cuerpo).status_code == 400


def test_uf_se_consulta_una_vez_y_queda_guardada(cliente_http, org_a, monkeypatch):
    llamadas = []

    def fuente_falsa(fecha):
        llamadas.append(fecha)
        return Decimal("39123.45")

    monkeypatch.setattr("app.indicadores.FUENTES", [("prueba", fuente_falsa)])
    h = login(cliente_http, org_a["email"])
    fecha = "2026-03-15"
    r1 = cliente_http.get("/api/indicadores/uf", params={"fecha": fecha}, headers=h)
    r2 = cliente_http.get("/api/indicadores/uf", params={"fecha": fecha}, headers=h)
    assert r1.status_code == r2.status_code == 200
    assert float(r1.json()["valor"]) == 39123.45
    assert len(llamadas) == 1  # la segunda vez sale de la base


def test_agenda_contactos_cuotas_y_recordatorios(cliente_http, org_a):
    h = login(cliente_http, org_a["email"])
    cartera = crear_cartera(cliente_http, h, n=41)
    hoy = hoy_chile()

    cliente_http.post("/api/gestiones/", headers=h, json={
        "cobranza_id": cartera["cobranza_id"], "descripcion": "Dice que paga el viernes",
        "fecha_proximo_contacto": hoy.isoformat()})
    r = cliente_http.post("/api/recordatorios", headers=h, json={
        "fecha": hoy.isoformat(), "hora": "10:30", "titulo": "Revisar transferencia",
        "cobranza_id": cartera["cobranza_id"]})
    assert r.status_code == 201
    recordatorio_id = r.json()["id"]
    cliente_http.post("/api/acuerdos/", headers=h, json={
        "cobranza_id": cartera["cobranza_id"], "monto_total_acordado": 90000, "numero_cuotas": 3,
        "fecha_primera_cuota": (hoy - timedelta(days=5)).isoformat()})

    items = cliente_http.get("/api/agenda/hoy", headers=h).json()
    tipos = {i["tipo"] for i in items}
    assert {"recordatorio", "cuota"} <= tipos
    assert any(i["tipo"] == "cuota" and i["atrasado"] for i in items)

    # La gestión del acuerdo (sin próximo contacto) dejó atendido el compromiso.
    assert not any(i["tipo"] == "contacto" for i in items)

    cliente_http.put(f"/api/recordatorios/{recordatorio_id}", headers=h, json={"estado": "hecho"})
    items = cliente_http.get("/api/agenda/hoy", headers=h).json()
    assert not any(i["tipo"] == "recordatorio" for i in items)

    # Un mes completo.
    mes = cliente_http.get("/api/agenda", headers=h, params={
        "desde": hoy.replace(day=1).isoformat(), "hasta": (hoy + timedelta(days=70)).isoformat()}).json()
    assert sum(1 for i in mes if i["tipo"] == "cuota") >= 2


def test_operador_solo_ve_su_agenda(cliente_http, org_a):
    ho = login(cliente_http, agregar_usuario(org_a["org_id"], "operador"))
    assert cliente_http.get("/api/agenda", headers=ho, params={"todos": True}).status_code == 403
    assert cliente_http.get("/api/agenda", headers=ho).status_code == 200


def test_mensaje_de_pago(cliente_http, org_a):
    h = login(cliente_http, org_a["email"])
    cartera = crear_cartera(cliente_http, h, monto=250000, n=42)
    r = cliente_http.get(f"/api/cobranzas/{cartera['cobranza_id']}/mensaje-pago", headers=h)
    assert r.status_code == 200
    m = r.json()
    assert "$250.000" in m["texto"] and "Transferencia" in m["texto"]
    assert m["whatsapp_url"].startswith("https://wa.me/56911112222?text=")
    assert m["falta_datos_pago"] is False

    cliente_http.put("/api/organizacion", headers=h, json={
        "plantilla_mensaje_pago": "Hola {nombre}, debes {saldo}. Paga en: {datos_pago}"})
    m = cliente_http.get(f"/api/cobranzas/{cartera['cobranza_id']}/mensaje-pago", headers=h).json()
    assert m["texto"].startswith("Hola Deudor, debes $250.000.")


def test_word_del_acuerdo(cliente_http, org_a):
    from io import BytesIO
    from docx import Document

    h = login(cliente_http, org_a["email"])
    cartera = crear_cartera(cliente_http, h, monto=600_000, n=43)
    r = cliente_http.post("/api/calculadora/acuerdo/crear", headers=h, json={
        "cobranza_id": cartera["cobranza_id"], "capital": 600000, "numero_cuotas": 3,
        "tasa_mensual": 1, "uf": 39000, "abono_inicial": 100000,
        "fecha_primera_cuota": (date.today() + timedelta(days=5)).isoformat()})
    acuerdo = r.json()
    r = cliente_http.get(f"/api/documentos/acuerdo/{acuerdo['id']}", headers=h)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    doc = Document(BytesIO(r.content))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "ACUERDO DE PAGO" in texto and "TOTAL PAGARÉ" in texto
    assert "3 cuotas iguales" in texto
    tabla = doc.tables[-1]
    assert tabla.rows[1].cells[0].text == "PIE"
    assert tabla.rows[-1].cells[-1].text.replace(".", "") == str(int(float(acuerdo["monto_total_acordado"])))


def test_panel_de_indicadores(cliente_http, org_a):
    h = login(cliente_http, org_a["email"])
    cartera = crear_cartera(cliente_http, h, monto=500_000, n=44)
    hoy = hoy_chile()
    cliente_http.post("/api/pagos/", headers=h, json={
        "cobranza_id": cartera["cobranza_id"], "monto": 120000, "capital": 100000,
        "honorarios": 20000, "fecha_pago": hoy.isoformat()})
    cliente_http.post("/api/acuerdos/", headers=h, json={
        "cobranza_id": cartera["cobranza_id"], "monto_total_acordado": 300000, "numero_cuotas": 3,
        "fecha_primera_cuota": (hoy - timedelta(days=40)).isoformat()})
    r = cliente_http.get("/api/panel", headers=h)
    assert r.status_code == 200, r.text
    p = r.json()
    assert float(p["recupero"]["capital"]) == 100000 and p["recupero"]["pagos"] == 1
    assert len(p["meses"]) == 12 and float(p["meses"][-1]["total"]) == 120000
    assert p["cuotas_atrasadas"] >= 1 and float(p["monto_atrasado"]) > 0
    assert any(float(c["recuperado_total"]) == 100000 for c in p["por_cliente"])
    assert sum(e["cantidad"] for e in p["cartera"]) == 1


def test_portal_de_mandantes(cliente_http, org_a):
    h = login(cliente_http, org_a["email"])
    propia = crear_cartera(cliente_http, h, monto=400_000, n=45)
    ajena = crear_cartera(cliente_http, h, monto=300_000, n=46)  # otro mandante, misma organización
    cliente_http.post("/api/gestiones/", headers=h, json={
        "cobranza_id": propia["cobranza_id"], "descripcion": "Llamada al deudor"})
    nota = next(t["id"] for t in cliente_http.get("/api/gestiones/tipos", headers=h).json() if t["codigo"] == "nota")
    cliente_http.post("/api/gestiones/", headers=h, json={
        "cobranza_id": propia["cobranza_id"], "tipo_id": nota, "descripcion": "NOTA INTERNA secreta"})
    acuerdo = cliente_http.post("/api/acuerdos/", headers=h, json={
        "cobranza_id": propia["cobranza_id"], "monto_total_acordado": 400000, "numero_cuotas": 2,
        "fecha_primera_cuota": (date.today() + timedelta(days=5)).isoformat()}).json()

    hm = login(cliente_http, agregar_usuario(org_a["org_id"], "mandante", cliente_id=propia["cliente_id"]))
    resumen = cliente_http.get("/api/portal/resumen", headers=hm).json()
    assert resumen["casos_totales"] == 1 and resumen["acuerdos_por_aprobar"] == 1

    lista = cliente_http.get("/api/portal/cobranzas", headers=hm).json()
    assert [c["id"] for c in lista] == [propia["cobranza_id"]]
    assert cliente_http.get(f"/api/portal/cobranzas/{ajena['cobranza_id']}", headers=hm).status_code == 404

    ficha = cliente_http.get(f"/api/portal/cobranzas/{propia['cobranza_id']}", headers=hm).json()
    textos = " ".join(g["descripcion"] for g in ficha["gestiones"])
    assert "Llamada al deudor" in textos and "NOTA INTERNA" not in textos

    assert cliente_http.post(f"/api/portal/acuerdos/{acuerdo['id']}/aprobar", headers=hm).status_code == 204
    assert cliente_http.get("/api/portal/acuerdos/pendientes", headers=hm).json() == []
    # El estudio ve la aprobación en el historial.
    historial = cliente_http.get("/api/gestiones/", params={"cobranza_id": propia["cobranza_id"]}, headers=h).json()
    assert any("APROBADO por el mandante" in g["descripcion"] for g in historial)

    # Un usuario interno no entra al portal, y el mandante no entra a la API interna.
    assert cliente_http.get("/api/portal/resumen", headers=h).status_code == 403
    assert cliente_http.get("/api/cobranzas/", headers=hm).status_code == 403


def test_logo_de_la_empresa(cliente_http, org_a, org_b):
    import io
    import zipfile
    import segno

    h = login(cliente_http, org_a["email"])
    cartera = crear_cartera(cliente_http, h, n=47)
    acuerdo = cliente_http.post("/api/acuerdos/", headers=h, json={
        "cobranza_id": cartera["cobranza_id"], "monto_total_acordado": 100000, "numero_cuotas": 2,
        "fecha_primera_cuota": (date.today() + timedelta(days=5)).isoformat()}).json()

    def archivos_word():
        r = cliente_http.get(f"/api/documentos/acuerdo/{acuerdo['id']}", headers=h)
        return zipfile.ZipFile(io.BytesIO(r.content))

    # Sin logo: va la esfera genérica del producto.
    from app.documento_acuerdo import LOGO_GENERICO
    z = archivos_word()
    imagenes = [z.read(n) for n in z.namelist() if n.startswith("word/media/")]
    assert LOGO_GENERICO.read_bytes() in imagenes

    # Un archivo que no es imagen se rechaza.
    r = cliente_http.put("/api/empresa/logo", headers=h,
                         files={"archivo": ("logo.png", b"<svg onload=alert(1)>", "image/png")})
    assert r.status_code == 422

    png = io.BytesIO()
    segno.make("logo").save(png, kind="png", scale=10)
    r = cliente_http.put("/api/empresa/logo", headers=h,
                         files={"archivo": ("logo.png", png.getvalue(), "image/png")})
    assert r.status_code == 200 and r.json()["tiene_logo"] is True
    r = cliente_http.get("/api/empresa/logo", headers=h)
    assert r.status_code == 200 and r.headers["content-type"] == "image/png"
    z = archivos_word()
    assert png.getvalue() in [z.read(n) for n in z.namelist() if n.startswith("word/media/")]

    # Cada organización ve solo su logo.
    hb = login(cliente_http, org_b["email"])
    assert cliente_http.get("/api/empresa/logo", headers=hb).status_code == 404
