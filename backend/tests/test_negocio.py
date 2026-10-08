"""
Reglas de negocio: acuerdos con cuotas, cascada de pagos, campos
personalizados, carga masiva y documentos.
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook

from tests.conftest import crear_cartera, login, rut_valido


def test_acuerdo_genera_cuotas_y_pagos_cierran_la_cobranza(cliente_http, org_a):
    h = login(cliente_http, org_a["email"])
    cartera = crear_cartera(cliente_http, h, monto=300000, n=20)

    r = cliente_http.post("/api/acuerdos/", headers=h, json={
        "cobranza_id": cartera["cobranza_id"], "monto_total_acordado": 300000,
        "numero_cuotas": 3, "fecha_primera_cuota": "2026-01-31", "capital": 300000,
    })
    assert r.status_code == 201, r.text
    acuerdo = r.json()
    assert [c["fecha_vencimiento"] for c in acuerdo["cuotas"]] == ["2026-01-31", "2026-02-28", "2026-03-31"]
    assert sum(float(c["monto"]) for c in acuerdo["cuotas"]) == 300000

    # Un solo acuerdo vigente.
    r = cliente_http.post("/api/acuerdos/", headers=h, json={
        "cobranza_id": cartera["cobranza_id"], "monto_total_acordado": 1,
        "numero_cuotas": 1, "fecha_primera_cuota": "2026-01-31",
    })
    assert r.status_code == 400

    for cuota in acuerdo["cuotas"]:
        r = cliente_http.post("/api/pagos/", headers=h, json={
            "cobranza_id": cartera["cobranza_id"], "cuota_id": cuota["id"],
            "monto": cuota["monto"], "capital": cuota["monto"], "forma_pago": "transferencia",
        })
        assert r.status_code == 201, r.text

    cob = cliente_http.get(f"/api/cobranzas/{cartera['cobranza_id']}", headers=h).json()
    assert cob["estado"] == "pagada" and float(cob["monto_actual"]) == 0

    historial = cliente_http.get("/api/gestiones/", params={"cobranza_id": cartera["cobranza_id"]},
                                 headers=h).json()
    textos = " | ".join(g["descripcion"] for g in historial)
    assert "ACUERDO DE PAGO" in textos and "CUENTA SALDADA" in textos


def test_solo_el_capital_descuenta_saldo(cliente_http, org_a):
    h = login(cliente_http, org_a["email"])
    cartera = crear_cartera(cliente_http, h, monto=100000, n=21)
    r = cliente_http.post("/api/pagos/", headers=h, json={
        "cobranza_id": cartera["cobranza_id"], "monto": 50000, "capital": 40000, "honorarios": 10000,
    })
    assert r.status_code == 201
    cob = cliente_http.get(f"/api/cobranzas/{cartera['cobranza_id']}", headers=h).json()
    assert float(cob["monto_actual"]) == 60000 and cob["estado"] == "activa"


def test_rut_invalido_se_rechaza(cliente_http, org_a):
    h = login(cliente_http, org_a["email"])
    r = cliente_http.post("/api/deudores/", headers=h, json={"rut": "12345678-9", "nombre": "X"})
    assert r.status_code == 422


def test_campos_personalizados(cliente_http, org_a):
    h = login(cliente_http, org_a["email"])
    cartera = crear_cartera(cliente_http, h, n=22)

    r = cliente_http.post("/api/campos", headers=h, json={
        "entidad": "cobranza", "etiqueta": "Previsión", "tipo": "seleccion",
        "opciones": ["FONASA", "ISAPRE"], "cliente_id": cartera["cliente_id"],
    })
    assert r.status_code == 201, r.text
    assert r.json()["clave"] == "prevision"
    cliente_http.post("/api/campos", headers=h, json={
        "entidad": "cobranza", "etiqueta": "Fecha de alta", "tipo": "fecha"})

    url = f"/api/cobranzas/{cartera['cobranza_id']}"
    assert cliente_http.put(url, headers=h, json={"datos_extra": {"prevision": "OTRA"}}).status_code == 422
    assert cliente_http.put(url, headers=h, json={"datos_extra": {"no_existe": 1}}).status_code == 422
    r = cliente_http.put(url, headers=h, json={"datos_extra": {"prevision": "FONASA", "fecha_de_alta": "2026-05-01"}})
    assert r.status_code == 200
    assert r.json()["datos_extra"] == {"prevision": "FONASA", "fecha_de_alta": "2026-05-01"}
    # Borrar un valor con null; el otro se conserva.
    r = cliente_http.put(url, headers=h, json={"datos_extra": {"prevision": None}})
    assert r.json()["datos_extra"] == {"fecha_de_alta": "2026-05-01"}


def test_etiquetas_por_organizacion(cliente_http, org_a, org_b):
    ha = login(cliente_http, org_a["email"])
    hb = login(cliente_http, org_b["email"])
    r = cliente_http.put("/api/organizacion", headers=ha, json={"etiquetas": {"cliente": "Mandante"}})
    assert r.status_code == 200 and r.json()["etiquetas"]["cliente"] == "Mandante"
    assert cliente_http.get("/api/organizacion", headers=hb).json()["etiquetas"]["cliente"] == "Cliente"
    assert cliente_http.put("/api/organizacion", headers=ha,
                            json={"etiquetas": {"inventada": "x"}}).status_code == 422


def test_tipos_de_gestion_propios(cliente_http, org_a, org_b):
    ha = login(cliente_http, org_a["email"])
    hb = login(cliente_http, org_b["email"])
    r = cliente_http.post("/api/gestiones/tipos", headers=ha, json={"nombre": "Visita notario", "categoria": "judicial"})
    assert r.status_code == 201
    nombres_a = {t["nombre"] for t in cliente_http.get("/api/gestiones/tipos", headers=ha).json()}
    nombres_b = {t["nombre"] for t in cliente_http.get("/api/gestiones/tipos", headers=hb).json()}
    assert "Visita notario" in nombres_a and "Visita notario" not in nombres_b
    assert "Abono" in nombres_a and "Abono" in nombres_b


def _excel(filas: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    for fila in filas:
        ws.append(fila)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_carga_masiva_por_titulos_y_campos_extra(cliente_http, org_a):
    h = login(cliente_http, org_a["email"])
    r = cliente_http.post("/api/clientes/", headers=h, json={"rut": rut_valido(30), "razon_social": "Clínica Sur"})
    cliente_http.post("/api/campos", headers=h, json={"entidad": "cobranza", "etiqueta": "Patente"})

    plantilla = load_workbook(BytesIO(cliente_http.get("/api/importar/plantilla", headers=h).content))
    titulos = [c.value for c in plantilla.active[1]]
    assert "Patente" in titulos and "RUT deudor*" in titulos

    contenido = _excel([
        ["Monto deuda", "Cliente", "Nombre deudor", "RUT deudor", "Patente", "Columna rara"],
        [150000, "Clínica Sur", "Pedro Soto", "12.345.678-5", "AB-CD-12", "x"],
        [99000, "Clínica Sur", "Rut Malo", "12345678-9", "", ""],
        [50000, "No Existe", "Ana", rut_valido(31), "", ""],
    ])
    r = cliente_http.post("/api/importar/cobranzas", headers=h,
                          files={"archivo": ("carga.xlsx", contenido,
                                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["cobranzas_creadas"] == 1 and len(res["errores"]) == 2
    assert res["columnas_ignoradas"] == ["Columna rara"]
    cob = cliente_http.get("/api/cobranzas/buscar", params={"q": "Pedro"}, headers=h).json()[0]
    assert cob["datos_extra"] == {"patente": "AB-CD-12"}


def test_exportaciones_y_documentos(cliente_http, org_a):
    h = login(cliente_http, org_a["email"])
    cartera = crear_cartera(cliente_http, h, n=23)
    cliente_http.post("/api/pagos/", headers=h, json={
        "cobranza_id": cartera["cobranza_id"], "monto": 1000, "capital": 1000,
        "fecha_pago": date.today().isoformat()})
    hoy = date.today()
    r = cliente_http.get("/api/exportar/recupero", params={"anio": hoy.year, "mes": hoy.month}, headers=h)
    assert r.status_code == 200
    filas = list(load_workbook(BytesIO(r.content)).active.iter_rows(values_only=True))
    assert len(filas) >= 2
    r = cliente_http.get(f"/api/documentos/estado-cuenta/{cartera['cobranza_id']}", headers=h)
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_auditoria_registra_cambios_con_organizacion(cliente_http, org_a, org_b):
    ha = login(cliente_http, org_a["email"])
    hb = login(cliente_http, org_b["email"])
    cartera = crear_cartera(cliente_http, ha, n=24)
    cliente_http.put(f"/api/cobranzas/{cartera['cobranza_id']}", headers=ha, json={"estado": "archivada"})
    log_a = cliente_http.get("/api/auditoria/", params={"tabla": "cobranzas"}, headers=ha).json()
    assert any(e["datos_nuevos"] and e["datos_nuevos"].get("estado") == "archivada" for e in log_a)
    assert cliente_http.get("/api/auditoria/", params={"tabla": "cobranzas"}, headers=hb).json() == []


def test_listados_muestran_deudor_y_resumen(cliente_http, org_a):
    h = login(cliente_http, org_a["email"])
    cartera = crear_cartera(cliente_http, h, monto=250000, n=25)
    lista = cliente_http.get("/api/cobranzas/", headers=h, params={"deudor_id": cartera["deudor_id"]}).json()
    assert len(lista) == 1 and lista[0]["deudor_nombre"] == "Deudor 25" and lista[0]["cliente_nombre"]
    encontrados = cliente_http.get("/api/deudores/buscar", headers=h, params={"q": "Deudor 25"}).json()
    assert encontrados[0]["cobranzas_abiertas"] == 1 and encontrados[0]["saldo_abierto"] == 250000
